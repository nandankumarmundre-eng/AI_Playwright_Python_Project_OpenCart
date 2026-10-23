from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from data.web.models import LoginCredentials, RegistrationDetails

REGISTERED_ACCOUNT_DATA_FILE = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "web"
    / "registered_account_data.xlsx"
)
_HEADERS = ("first_name", "last_name", "email", "telephone", "password")


def save_registered_account(
    details: RegistrationDetails,
    file_path: Path = REGISTERED_ACCOUNT_DATA_FILE,
) -> None:
    if file_path.exists():
        try:
            workbook = load_workbook(file_path)
        except (OSError, InvalidFileException) as error:
            raise ValueError(
                f"Could not open registered-account workbook '{file_path}'."
            ) from error
        worksheet = workbook.active
        if worksheet is None:
            workbook.close()
            raise ValueError(
                f"Registered-account workbook '{file_path}' has no worksheet."
            )
        if worksheet.max_row == 1 and all(
            cell.value is None for cell in worksheet[1]
        ):
            worksheet.append(_HEADERS)
        elif tuple(cell.value for cell in worksheet[1]) != _HEADERS:
            workbook.close()
            raise ValueError(
                f"Registered-account workbook '{file_path}' has unexpected "
                "column headers."
            )
    else:
        workbook = Workbook()
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("Could not create a registered-account worksheet.")
        worksheet.title = "accounts"
        worksheet.append(_HEADERS)

    worksheet.append(
        (
            details.first_name,
            details.last_name,
            details.email,
            details.telephone,
            details.password,
        )
    )
    try:
        workbook.save(file_path)
    finally:
        workbook.close()


def read_latest_login_credentials(
    file_path: Path = REGISTERED_ACCOUNT_DATA_FILE,
) -> LoginCredentials:
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except (OSError, InvalidFileException) as error:
        raise ValueError(
            f"Could not open registered-account workbook '{file_path}'."
        ) from error

    try:
        worksheet = workbook.active
        if worksheet is None or tuple(
            cell.value for cell in next(worksheet.iter_rows(max_row=1))
        ) != _HEADERS:
            raise ValueError(
                f"Registered-account workbook '{file_path}' has invalid headers."
            )
        values = next(
            worksheet.iter_rows(
                min_row=worksheet.max_row,
                max_row=worksheet.max_row,
                values_only=True,
            ),
            None,
        )
        if values is None or len(values) != len(_HEADERS) or not all(values):
            raise ValueError(
                f"Registered-account workbook '{file_path}' has no complete "
                "account record."
            )
        return LoginCredentials(email=str(values[2]), password=str(values[4]))
    finally:
        workbook.close()
