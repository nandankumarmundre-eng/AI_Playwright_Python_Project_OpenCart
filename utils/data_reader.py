import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

WEB_DATA_DIRECTORY = Path(__file__).resolve().parents[1] / "data" / "web"


class DataReaderError(ValueError):
    """Raised when a test-data file is invalid or incomplete."""


def read_excel_data(
    file_name: str, sheet_name: str
) -> list[dict[str, str]]:
    workbook_path = WEB_DATA_DIRECTORY / file_name
    try:
        workbook = load_workbook(
            workbook_path, read_only=True, data_only=True
        )
    except (OSError, InvalidFileException) as error:
        raise DataReaderError(
            f"Could not open Excel test data '{workbook_path}'."
        ) from error

    try:
        if sheet_name not in workbook.sheetnames:
            raise DataReaderError(
                f"Worksheet '{sheet_name}' was not found in '{workbook_path}'."
            )
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise DataReaderError(
                f"Worksheet '{sheet_name}' in '{workbook_path}' is empty."
            )

        normalized_headers = [
            str(value).strip() if value is not None else ""
            for value in headers
        ]
        populated_headers = [header for header in normalized_headers if header]
        if not populated_headers:
            raise DataReaderError(
                f"Worksheet '{sheet_name}' in '{workbook_path}' has no headers."
            )
        if len(set(populated_headers)) != len(populated_headers):
            raise DataReaderError(
                f"Worksheet '{sheet_name}' in '{workbook_path}' "
                "contains duplicate headers."
            )

        records: list[dict[str, str]] = []
        for values in rows:
            record = {
                header: (
                    ""
                    if index >= len(values) or values[index] is None
                    else str(values[index]).strip()
                )
                for index, header in enumerate(normalized_headers)
                if header
            }
            if any(record.values()):
                records.append(record)
        return records
    finally:
        workbook.close()


def read_json(path: Path) -> Any:
    try:
        with path.open(encoding="utf-8") as file:
            return json.load(file)
    except (OSError, json.JSONDecodeError) as error:
        raise DataReaderError(f"Could not read JSON test data '{path}'.") from error


def read_api_json(file_name: str) -> Any:
    api_data_directory = WEB_DATA_DIRECTORY.parent / "api"
    return read_json(api_data_directory / file_name)
